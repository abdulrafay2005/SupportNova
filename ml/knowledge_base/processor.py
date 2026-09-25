"""
SupportNova Help Center Document Processor.

Supports:
- TXT
- DOCX
- PDF

Documents are extracted into a normalized structure that can
later be validated and stored in the SupportNova Help Center.
"""

from pathlib import Path


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".docx",
    ".pdf",
}


class DocumentProcessingError(Exception):
    """Raised when a document cannot be processed."""
    pass


def validate_file_path(file_path):
    """
    Validate that the supplied path points to a supported file.
    """

    path = Path(file_path)

    if not path.exists():
        raise DocumentProcessingError(
            f"File does not exist: {path}"
        )

    if not path.is_file():
        raise DocumentProcessingError(
            f"Path is not a file: {path}"
        )

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError(
            f"Unsupported file type: {path.suffix}"
        )

    return path


def extract_txt(file_path):
    """
    Extract text from a TXT file.
    """

    path = validate_file_path(file_path)

    try:
        return path.read_text(
            encoding="utf-8"
        ).strip()

    except UnicodeDecodeError:
        try:
            return path.read_text(
                encoding="utf-8-sig"
            ).strip()

        except Exception as error:
            raise DocumentProcessingError(
                f"Could not read TXT file: {error}"
            )


def extract_docx(file_path):
    """
    Extract paragraphs and table contents from a DOCX file.
    """

    path = validate_file_path(file_path)

    try:
        from docx import Document
    except ImportError:
        raise DocumentProcessingError(
            "python-docx is not installed. "
            "Run: pip install python-docx"
        )

    try:
        document = Document(path)

        parts = []

        # Paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts).strip()

    except Exception as error:
        raise DocumentProcessingError(
            f"Could not read DOCX file: {error}"
        )


def extract_pdf(file_path):
    """
    Extract text from a PDF file.
    """

    path = validate_file_path(file_path)

    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocumentProcessingError(
            "pypdf is not installed. "
            "Run: pip install pypdf"
        )

    try:
        reader = PdfReader(path)

        pages = []

        for page in reader.pages:

            text = page.extract_text()

            if text:
                pages.append(text.strip())

        return "\n".join(pages).strip()

    except Exception as error:
        raise DocumentProcessingError(
            f"Could not read PDF file: {error}"
        )


def extract_text(file_path):
    """
    Extract text based on file extension.
    """

    path = validate_file_path(file_path)

    extension = path.suffix.lower()

    if extension == ".txt":
        return extract_txt(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension == ".pdf":
        return extract_pdf(path)

    raise DocumentProcessingError(
        f"Unsupported extension: {extension}"
    )


def process_document(file_path):
    """
    Process one document and return normalized metadata.
    """

    path = validate_file_path(file_path)

    text = extract_text(path)

    return {
        "filename": path.name,
        "file_path": str(path),
        "extension": path.suffix.lower(),
        "text": text,
        "character_count": len(text),
        "word_count": len(text.split()),
    }


def process_directory(directory_path):
    """
    Recursively process all supported documents inside a directory.

    Subdirectories such as:
        policies/
        sops/
        rules/

    are included automatically.
    """

    directory = Path(directory_path)

    if not directory.exists():
        raise DocumentProcessingError(
            f"Directory does not exist: {directory}"
        )

    if not directory.is_dir():
        raise DocumentProcessingError(
            f"Path is not a directory: {directory}"
        )

    documents = []

    # IMPORTANT:
    # rglob() recursively searches policies/, sops/, rules/, etc.
    for file_path in sorted(
        directory.rglob("*")
    ):

        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:

            document = process_document(
                file_path
            )

            documents.append(
                document
            )

        except DocumentProcessingError as error:

            print(
                f"[ERROR] {file_path.name}: {error}"
            )

    return documents